from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path('/home/abhijit-dash/Desktop/Campus AI Project/Campus_Drive_AI_Project_Demonstration.docx')
NAVY = '17365D'
BLUE = '1F4E79'
TEAL = '087E8B'
PALE = 'EAF2F8'
LIGHT = 'F3F6F9'
DARK = '263238'
MUTED = '586575'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False, color=DARK, size=9.2):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = 'Aptos'
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    set_repeat_table_header(t.rows[0])
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, 'FFFFFF', 9)
        shade(t.rows[0].cells[i], NAVY)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], str(val), False, DARK, 8.7)
            if ri % 2 == 1:
                shade(cells[i], LIGHT)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p

def add_number(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p

def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), 'F1F4F7'); pPr.append(shd)
    r = p.add_run(text)
    r.font.name = 'Consolas'; r.font.size = Pt(8.3); r.font.color.rgb = RGBColor.from_string('263238')
    return p

def add_note(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    lead = p.add_run(label + ': ')
    lead.bold = True
    lead.font.color.rgb = RGBColor.from_string(TEAL)
    p.add_run(text)
    return p

def add_page_field(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run('Campus Drive AI  |  Page ')
    r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(MUTED)
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
    p._p.append(fld)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.68); sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.78); sec.right_margin = Inches(0.78)
sec.header_distance = Inches(0.3); sec.footer_distance = Inches(0.3)
styles = doc.styles
normal = styles['Normal']; normal.font.name='Aptos'; normal.font.size=Pt(9.5); normal.font.color.rgb=RGBColor.from_string(DARK)
normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.08
for name, size, color in [('Title', 27, NAVY), ('Heading 1', 17, NAVY), ('Heading 2', 12, TEAL), ('Heading 3', 10, BLUE)]:
    st=styles[name]; st.font.name='Aptos Display' if name != 'Heading 3' else 'Aptos'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
    st.paragraph_format.space_before=Pt(10 if name!='Title' else 0); st.paragraph_format.space_after=Pt(5)
    st.paragraph_format.keep_with_next=True
title_ppr = styles['Title']._element.pPr
if title_ppr is not None:
    borders = title_ppr.find(qn('w:pBdr'))
    if borders is not None:
        title_ppr.remove(borders)
# header/footer
hp=sec.header.paragraphs[0]; hp.text='CAMPUS DRIVE AI  |  ROS 2 AND GAZEBO ROBOTICS'; hp.alignment=WD_ALIGN_PARAGRAPH.LEFT
for r in hp.runs: r.font.name='Aptos'; r.font.size=Pt(8); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL)
add_page_field(sec.footer.paragraphs[0])

# Cover / executive summary
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(42); p.paragraph_format.space_after=Pt(8)
r=p.add_run('PROJECT DEMONSTRATION REPORT'); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(TEAL)
p=doc.add_paragraph(style='Title'); p.add_run('Campus Drive AI')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(18)
r=p.add_run('ROS 2 and Gazebo campus robot simulation'); r.font.size=Pt(15); r.font.color.rgb=RGBColor.from_string(MUTED)
add_table(doc, ['Student', 'Programme', 'Institution', 'Demonstration'], [[
    'Abhijit Dash', 'B.Tech Computer Science and Engineering, 4th Year', 'Centurion University, Bhubaneswar', '16 October 2026'
]], [1.45, 2.0, 1.9, 1.4])
doc.add_heading('Project summary', level=1)
doc.add_paragraph('Campus Drive AI is a learning-focused robotics simulation in which a small differential-drive robot travels through a campus courtyard. ROS 2 Jazzy runs the robot logic, while Gazebo Harmonic provides the 3D world and simulated motion. The robot reads 2D LiDAR and wheel odometry, follows nine programmed waypoints along an S-shaped route around three blocks, reacts to nearby obstacles, and records performance data.')
doc.add_paragraph('The current project is a rule-based robotics prototype. It demonstrates sensing, control, simulation, and evaluation; it does not use a trained machine-learning model, SLAM, or the Nav2 navigation stack.')
add_note(doc, 'Demonstration goal', 'Show the campus scene, the car-style robot, LiDAR and odometry topics, waypoint progress, obstacle response, and the completion indicator that changes from SIMULATION STARTED to ROUTE FINISHED.')
doc.add_heading('Project flow at a glance', level=2)
add_table(doc, ['Stage', 'What happens'], [
    ('1. Simulate', 'Gazebo loads the campus world and spawns the robot.'),
    ('2. Sense', 'The robot publishes LiDAR scans on /scan and odometry on /odom.'),
    ('3. Decide', 'The ROS 2 controller tracks the next waypoint and checks the forward LiDAR sector.'),
    ('4. Move', 'The controller publishes speed commands on /cmd_vel; a bridge passes them to Gazebo.'),
    ('5. Evaluate', 'The evaluator records motion and sensor metrics in a CSV file.')
], [1.35, 5.4])

doc.add_page_break()
# Core report
doc.add_heading('1. Introduction', level=1)
doc.add_paragraph('Mobile robots can help with campus tasks such as carrying items, guiding visitors, or inspecting walkways. Before testing a robot in a real environment, developers can use simulation to study its movement and sensor data safely. This project builds a small campus-like environment and demonstrates how a ROS 2 robot can move through it using a simple navigation loop.')
doc.add_heading('Problem statement', level=2)
doc.add_paragraph('A robot moving through a campus must estimate where it is, identify nearby obstacles, choose a direction, and reach its destinations without colliding. The project provides a repeatable simulated setting for observing those functions and measuring a run.')
doc.add_heading('2. Objectives', level=1)
for x in [
    'Create a campus courtyard world in Gazebo with buildings, paved paths, lawns, trees, and planters.',
    'Model a differential-drive robot with a vehicle-style body, LiDAR, and wheel odometry.',
    'Connect Gazebo sensor and control messages to ROS 2 topics.',
    'Follow nine S-shaped route targets from the green start pad at (-3.8, 1.7) m to the red destination at (3.8, 1.7) m, passing around three fixed blocks.',
    'Use forward LiDAR distance and side clearance to trigger a reactive obstacle response.',
    'Display route status in the simulation and save evaluation measurements to CSV.'
]: add_bullet(doc, x)
doc.add_heading('Project scope', level=2)
doc.add_paragraph('The demonstration focuses on the complete simulation loop: a world model, robot sensors, a ROS 2 controller, a Gazebo command bridge, and a measurement recorder. The controlled environment makes it possible to inspect robot behavior and repeat runs without putting a physical platform or people at risk.')
doc.add_paragraph('The project is designed as a foundation for campus robotics learning. It shows how separate components communicate and how a local sensor rule can change the robot’s motion. Its scope is intentionally smaller than a production delivery robot, which would require mapping, localization, robust path planning, and extensive safety validation.')
doc.add_page_break()
doc.add_heading('3. Technology and software stack', level=1)
add_table(doc, ['Technology', 'Role in the project'], [
    ('Ubuntu 24.04', 'Operating system used to run the development environment.'),
    ('ROS 2 Jazzy', 'Robot software framework; runs the controller, evaluator, and ROS topics.'),
    ('Gazebo Harmonic / Gazebo Sim 8', 'Physics-based 3D simulation of the campus, robot, and sensors.'),
    ('SDF', 'Describes the Gazebo world, robot model, sensors, and simulation systems.'),
    ('Python 3 and rclpy', 'Implements the waypoint, LiDAR-avoidance, status, and evaluation nodes.'),
    ('ros_gz_bridge', 'Connects Gazebo messages with ROS 2 topics and services.'),
    ('colcon and CMake / ament', 'Build and install the ROS 2 workspace packages.')
], [2.15, 4.6])
doc.add_heading('Main ROS 2 topics', level=2)
add_table(doc, ['Topic', 'Message type', 'Purpose'], [
    ('/scan', 'sensor_msgs/msg/LaserScan', 'LiDAR distance measurements.'),
    ('/odom', 'nav_msgs/msg/Odometry', 'Estimated robot position and velocity.'),
    ('/cmd_vel', 'geometry_msgs/msg/Twist', 'Linear and angular velocity commands.'),
    ('/campus_navigation/status', 'std_msgs/msg/String', 'Current waypoint and distance.'),
    ('/campus_navigation/obstacle_detected', 'std_msgs/msg/Bool', 'Latches true after a valid nearby LiDAR return.'),
    ('/campus_navigation/succeeded', 'std_msgs/msg/Bool', 'Signals when all route targets are reached.'),
    ('/campus_navigation/mode', 'std_msgs/msg/String', 'Selects automatic or manual control.')
], [2.05, 2.25, 2.45])
doc.add_heading('How the software components work together', level=2)
doc.add_paragraph('The launch file starts the simulator and the ROS 2 processes as one demonstration. The bridge translates messages between the Gazebo Transport system and ROS 2. The autonomy node consumes sensor and pose data and publishes velocity commands. The evaluator listens to the same run and stores measurements separately, so the control logic and reporting logic remain easy to understand.')
doc.add_paragraph('A topic is a named stream of messages. For example, /scan is continuously updated by the simulated LiDAR, while /cmd_vel carries the controller’s requested motion. A second terminal can inspect these streams while the launch terminal continues running the simulation.')
doc.add_page_break()
doc.add_heading('4. Methodology and project flow', level=1)
doc.add_paragraph('The simulation uses a closed loop. The robot senses the world, the ROS 2 controller selects motion, and Gazebo applies the command. New sensor data then updates the controller.')
for h, para in [
    ('4.1 Campus world and robot model', 'The SDF world contains campus buildings, a paved courtyard, lawn areas, trees, planters, and three fixed blocks. The robot has a differential-drive base, a blue vehicle-style shell, a planar LiDAR, and wheel odometry. The green start pad is at (-3.8, 1.7) m; the separate red destination is at (3.8, 1.7) m, 7.6 m away.'),
    ('4.2 Sensor and command bridge', 'The ros_gz bridge makes Gazebo LiDAR data available as /scan and odometry available as /odom. Velocity commands published to /cmd_vel are bridged back to Gazebo so its differential-drive system can move the robot.'),
    ('4.3 Waypoint navigation', 'The autonomy node stores nine world-frame target points that form an S-shaped route around the blocks. It captures spawn odometry as the local origin and converts each target to that frame. At each control step it calculates distance and heading, turns toward the target, and drives forward when aligned. It advances within 0.12 m of a target and stops after a 180-second watchdog timeout.'),
    ('4.4 Reactive obstacle avoidance', 'The controller checks the forward LiDAR sector and compares left and right clearance. Below the 0.65 m avoidance threshold it turns toward the clearer side, or stops if neither side has enough room. Returns shorter than 0.40 m are filtered by the controller to reduce self-echoes. This reactive layer complements the fixed roadmap; it does not create a map or provide global path planning.'),
    ('4.5 Status and interactive operation', 'The display node spawns a SIMULATION STARTED sign when Gazebo is ready, then replaces it with ROUTE FINISHED after route success. Gazebo provides pause, play, and camera controls. A separate keyboard node can request manual driving; W/I move forward, S/K reverse, A/J turn left, D/L turn right, Space stops, and Q requests automatic mode again.'),
    ('4.6 Evaluation', 'The evaluator listens to odometry, LiDAR, velocity commands, and route completion. It writes a CSV containing elapsed time, odometry-based path length, minimum raw LiDAR returns, scan counts, zero-velocity command samples, and route success. The CSV is saved on successful completion or when the evaluator shuts down.')
]:
    if h.startswith('4.3'):
        doc.add_page_break()
    doc.add_heading(h, level=2); doc.add_paragraph(para)
    if h.startswith('4.2'):
        doc.add_heading('Campus scene and model components', level=2)
        doc.add_paragraph('The environment is built from simple geometric shapes and colors so that the layout is light enough to simulate on a student laptop. Static scenery defines the campus boundaries; the robot and its sensor links remain dynamic.')
        add_table(doc, ['Element', 'Purpose in the scene'], [
            ('Yellow S-shaped route', 'Guides the robot below block 1, above block 2, and below block 3.'),
            ('Campus buildings', 'Give the courtyard a recognizable campus setting; the purple building carries the university name.'),
            ('Trees and planters', 'Provide LiDAR-visible objects for nearby obstacle observations.'),
            ('Start and finish markers', 'Mark separate endpoints at (-3.8, 1.7) m and (3.8, 1.7) m.'),
            ('Vehicle-style robot shell', 'Makes the differential-drive platform easy to identify during the demo.')
        ], [2.35, 4.4])
doc.add_heading('System flow', level=2)
add_table(doc, ['Input', 'Processing', 'Output'], [
    ('/scan and /odom', 'Waypoint logic plus forward-sector obstacle check', '/cmd_vel'),
    ('Gazebo world and robot model', 'Physics and simulated sensor systems', 'Robot movement and sensor messages'),
    ('/scan, /odom, /cmd_vel, route result', 'Evaluation recorder', 'results/campus_evaluation.csv')
], [2.0, 2.7, 2.05])
doc.add_heading('Programmed route reference', level=2)
doc.add_paragraph('The controller uses these nine targets in order. Coordinates are in meters in the Gazebo world frame. The robot spawns at (-3.8, 1.7) m and ends at the separate red pad at (3.8, 1.7) m. The three blocks are centered at (-2.5, 1.7), (0, 1.7), and (2.5, 1.7) m.')
add_table(doc, ['Waypoint', 'Target (x, y)', 'Purpose'], [
    ('1', '(-3.2, 1.70)', 'Line up before block 1.'),
    ('2', '(-3.2, 0.65)', 'Pass below block 1.'),
    ('3', '(-1.35, 0.65)', 'Travel between blocks 1 and 2.'),
    ('4', '(-1.35, 2.75)', 'Turn around the left side of block 2.'),
    ('5', '(1.35, 2.75)', 'Pass above block 2.'),
    ('6', '(1.35, 0.65)', 'Turn below block 2.'),
    ('7', '(3.2, 0.65)', 'Pass below block 3.'),
    ('8', '(3.2, 1.70)', 'Rejoin the east-side path.'),
    ('9', '(3.8, 1.70)', 'Reach the red destination pad.')
], [1.2, 1.6, 3.95])
# The source tree is described in the README; omit a standalone inventory page.
doc.add_heading('5. Procedure for the demonstration', level=1)
doc.add_heading('Before starting', level=2)
for x in [
    'Use the Ubuntu 24.04 installation where ROS 2 Jazzy and Gazebo Harmonic are available.',
    'Open a terminal and make sure the project folder is ~/Desktop/campus_ws.',
    'Close any earlier Gazebo or ROS launch session so that only one simulation is running.'
]: add_bullet(doc, x)
doc.add_heading('Confirm the software and workspace', level=2)
doc.add_paragraph('Run this quick check before the demonstration. It confirms that ROS 2 is sourced, Gazebo responds, and colcon can see the two project packages:')
add_code(doc, 'source /opt/ros/jazzy/setup.bash && ros2 pkg prefix rclcpp && gz sim --version && cd ~/Desktop/campus_ws && colcon list')
add_bullet(doc, 'Expected: the ROS package prefix is /opt/ros/jazzy, Gazebo reports version 8.x, and colcon lists campus_robot_bringup and campus_robot_description.')
add_bullet(doc, 'Replace 123 in the second-terminal commands with the numeric ROS_DOMAIN_ID printed by the launch terminal.')
doc.add_heading('Build and launch', level=2)
doc.add_paragraph('Copy and run this command in one terminal. It loads the ROS 2 environment, builds both workspace packages, and starts the graphical simulation:')
add_code(doc, 'cd ~/Desktop/campus_ws && source /opt/ros/jazzy/setup.bash && colcon build --symlink-install && source install/setup.bash && ros2 launch campus_robot_bringup campus_sim.launch.py')
doc.add_paragraph('Keep this terminal open. Wait for Gazebo to open and for the robot to appear. The route starts automatically. Use the Gazebo pause/play controls or mouse camera controls to present the scene.')
doc.add_heading('What to look for when Gazebo opens', level=2)
add_bullet(doc, 'Check that the campus courtyard, paved route, buildings, trees, and planters are visible.')
add_bullet(doc, 'Locate the blue vehicle-style robot at the green start pad (-3.8, 1.7 m) and the separate red destination pad (3.8, 1.7 m).')
add_bullet(doc, 'Find the CENTURION UNIVERSITY lettering on the front-facing wall of the purple building.')
add_bullet(doc, 'The autonomous route begins after odometry and LiDAR are available. Keep the launch terminal open to see waypoint, obstacle, and completion messages.')
doc.add_page_break()
doc.add_heading('Verify that ROS 2 topics are active', level=2)
doc.add_paragraph('In a second terminal, list the topics. This confirms that the launch created the bridge and navigation nodes:')
add_code(doc, 'source /opt/ros/jazzy/setup.bash && source ~/Desktop/campus_ws/install/setup.bash && export ROS_DOMAIN_ID=123 && ros2 topic list')
doc.add_paragraph('Set ROS_DOMAIN_ID to the number printed by the launch terminal (replace `123` in the command), then look for /scan, /odom, /cmd_vel, /campus_navigation/status, /campus_navigation/obstacle_detected, and /campus_navigation/succeeded. If ros2 is not found, source the ROS and workspace setup files in that terminal. If topics are absent, check the launch terminal and the domain value.')
doc.add_heading('Inspect route and sensor messages', level=2)
doc.add_paragraph('Use --once to print one message and return to the prompt. Run one command at a time in the second terminal:')
add_code(doc, 'source /opt/ros/jazzy/setup.bash && source ~/Desktop/campus_ws/install/setup.bash && export ROS_DOMAIN_ID=123 && ros2 topic echo --once /campus_navigation/status')
add_code(doc, 'source /opt/ros/jazzy/setup.bash && source ~/Desktop/campus_ws/install/setup.bash && export ROS_DOMAIN_ID=123 && ros2 topic echo --once /scan')
add_code(doc, 'source /opt/ros/jazzy/setup.bash && source ~/Desktop/campus_ws/install/setup.bash && export ROS_DOMAIN_ID=123 && ros2 topic echo --once /odom')
doc.add_paragraph('Explain the messages simply: /scan contains LiDAR distances, /odom contains the robot pose and velocity estimate, and /campus_navigation/status tells which waypoint is active.')
doc.add_heading('Optional manual driving', level=2)
doc.add_paragraph('If time permits, start the keyboard controller in a third terminal. This is optional; keep the autonomous route in control for the completion demonstration:')
add_code(doc, 'source /opt/ros/jazzy/setup.bash && source ~/Desktop/campus_ws/install/setup.bash && export ROS_DOMAIN_ID=123 && ros2 run campus_robot_bringup campus_teleop.py')
doc.add_paragraph('W/I move forward, S/K reverse, A/J turn left, D/L turn right, Space stops, and Q requests autonomous mode again. Use the terminal running teleoperation to enter keys.')
doc.add_heading('Complete and stop the run', level=2)
for x in [
    'Wait for the controller to reach waypoint 9 and announce that the route is complete.',
    'Confirm the in-world display says ROUTE FINISHED and the launch terminal shows the evaluator saved the CSV.',
    'After showing the completed route and confirming the CSV was saved, press Ctrl+C once in the launch terminal to shut down Gazebo and the ROS nodes.'
]: add_number(doc, x)
add_note(doc, 'Keep the run clean', 'Do not launch a second Gazebo GUI against the existing world. If a previous simulation is still active, stop it with Ctrl+C before beginning this procedure.')
doc.add_heading('Quick troubleshooting', level=2)
add_table(doc, ['Symptom', 'What to do'], [
    ('ros2: command not found', 'Source /opt/ros/jazzy/setup.bash and then the workspace install/setup.bash in that terminal.'),
    ('ament_cmake cannot be found', 'Source the ROS 2 setup file before running colcon build.'),
    ('Gazebo reports an existing sun', 'Stop the duplicate GUI or launch with Ctrl+C; run only the project launch command.'),
    ('No robot or topics appear', 'Check launch output and confirm only one simulation is active.')
], [2.25, 4.45])

doc.add_heading('6. Expected results and evaluation', level=1)
doc.add_paragraph('During a successful demonstration, Gazebo should show the campus environment and vehicle-style robot. ROS 2 should publish LiDAR data on /scan, odometry on /odom, and velocity commands on /cmd_vel. The controller should publish waypoint status, respond to nearby obstacles, reach the red destination after nine route targets, and trigger the ROUTE FINISHED sign. The evaluator should save a CSV file.')
doc.add_heading('Evaluation measures', level=2)
add_table(doc, ['Metric', 'Latest saved successful run'], [
    ('Route result', 'Succeeded (route_succeeded=True)'),
    ('Elapsed time', '137.96 s'),
    ('Odometry path length', '13.058 m'),
    ('Minimum raw LiDAR return', '0.260 m'),
    ('Minimum raw forward return (±30°)', '0.581 m'),
    ('LiDAR scans', '1,117'),
    ('Scans with raw nearest range under 1 m', '1,022'),
    ('Zero-velocity command samples', '10')
], [2.3, 4.45])
add_note(doc, 'Latest saved successful run', 'The CSV records route_succeeded=True, 137.96 s elapsed, 13.058 m odometry path, 0.260 m minimum raw LiDAR return, 0.581 m minimum raw forward return, 1,117 scans, 1,022 scans with a raw return under 1 m, and 10 zero-velocity command samples. These are single-run values. Raw minima may include robot self-echoes and are not verified collision-clearance measurements. Three summary graphs in graphs/ show all available result points from the saved and verified evaluation CSV files.')
doc.add_heading('What counts as a successful run', level=2)
add_table(doc, ['Check', 'Evidence to show'], [
    ('World and robot start', 'Gazebo shows the courtyard and the robot spawn.'),
    ('Sensing', 'A /scan message contains LiDAR ranges; /odom contains a pose.'),
    ('Navigation', 'Waypoint status advances and /cmd_vel carries motion commands.'),
    ('Completion', 'The controller reports route completion and the sign changes to ROUTE FINISHED.'),
    ('Evaluation', 'results/campus_evaluation.csv is saved with route_succeeded true.')
], [1.8, 4.95])
doc.add_paragraph('When reviewing the CSV, first check route_succeeded. Compare path length and elapsed time between runs only when the world and controller settings are unchanged. The minimum LiDAR distance describes the closest detected return; it is not a complete proof that every possible collision was avoided.')
doc.add_page_break()
doc.add_heading('7. Limitations', level=1)
for x in [
    'Obstacle avoidance uses local LiDAR rules; it does not build a map or plan around obstacles globally.',
    'The project does not currently use SLAM, map-based localization, Nav2, or a trained AI/ML model.',
    'Results can change with simulation speed, controller tuning, and the map. Current-map performance should be verified by running the route.',
    'This is a simulation prototype and is not validated for controlling a real robot.'
]: add_bullet(doc, x)
doc.add_heading('8. Future enhancements', level=1)
for x in [
    'Add SLAM and map-based localization so the robot can build and use a campus map.',
    'Integrate Nav2 for global and local path planning, recovery behaviors, and goal management.',
    'Improve obstacle avoidance with dynamic obstacle detection and safer speed control.',
    'Add a camera and image-processing tasks such as sign or landmark recognition.',
    'Create repeatable test scenarios and compare route time, path length, clearance, and success rate over multiple runs.'
]: add_bullet(doc, x)
doc.add_heading('Suggested development order', level=2)
add_table(doc, ['Stage', 'Next work', 'Benefit'], [
    ('1. Reliability', 'Repeat runs and tune the waypoint and obstacle thresholds.', 'Establish dependable baseline behavior.'),
    ('2. Mapping', 'Add SLAM and map-based localization.', 'Let the robot reason about the wider campus.'),
    ('3. Planning', 'Integrate Nav2 and test global/local planners.', 'Navigate to goals with recovery behaviors.'),
    ('4. Perception', 'Add a camera and landmark/sign recognition.', 'Use visual information with LiDAR.'),
    ('5. Evaluation', 'Add contact monitoring and repeated controlled trials.', 'Measure success and collisions across runs.')
], [1.15, 3.2, 2.4])
doc.add_heading('Expected long-term outcome', level=2)
doc.add_paragraph('With these stages, the project can grow from a single scripted route into a testable campus navigation platform. Each upgrade should be measured against repeatable scenarios, with safety checks and route-completion evidence recorded alongside performance numbers.')
doc.add_heading('9. Conclusion', level=1)
doc.add_paragraph('Campus Drive AI brings together a simulated campus, a differential-drive robot, LiDAR and odometry, ROS 2 communication, a nine-target S-shaped route around three fixed blocks, reactive obstacle handling, an in-world completion sign, and automatic performance logging. The latest saved successful run completed in 137.96 seconds with a 13.058 m odometry path. Its raw LiDAR metrics require caution because self-echoes are not excluded by the evaluator. The implementation is a learning prototype that can be extended with repeat trials, collision monitoring, mapping, and more advanced navigation.')

# Document metadata and save
doc.core_properties.title = 'Campus Drive AI Project Demonstration Report'
doc.core_properties.subject = 'ROS 2 Jazzy and Gazebo Harmonic campus robotics simulation'
doc.core_properties.author = 'Abhijit Dash'
doc.core_properties.keywords = 'Campus Drive AI, ROS 2, Gazebo, robotics, LiDAR, navigation'
doc.save(OUT)
print(OUT)
